from pathlib import Path
import json
import re

from lxml import html
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "资料收集" / "网站群建设资料收集说明.docx"
AUDIT = Path(__file__).with_name("current_site_audit.json")
FULL_NAME = "厦门大学闽台非遗文化数字化保护与智能处理文化和旅游部重点实验室"


def page_soup(path):
    return html.fromstring((ROOT / path).read_text(encoding="utf-8"))


def text_content(element):
    return element.text_content().strip()


def audit_site():
    records = {}
    for folder in ("news", "notices"):
        items = []
        for path in sorted((ROOT / folder).glob("*.html"), reverse=True):
            soup = page_soup(path.relative_to(ROOT))
            content_matches = soup.xpath('//*[@class="original-content"]')
            content = content_matches[0] if content_matches else None
            meta_matches = soup.xpath('//*[@class="article-meta"]')
            meta = meta_matches[0] if meta_matches else None
            items.append({
                "path": str(path.relative_to(ROOT)),
                "title": text_content(soup.xpath('//h1')[0]),
                "meta": text_content(meta) if meta is not None else "",
                "body_characters": len(text_content(content)) if content is not None else 0,
                "images": content.xpath('.//img/@src') if content is not None else [],
                "links": content.xpath('.//a/@href') if content is not None else [],
            })
        records[folder] = items
    people = page_soup("people.html")
    records["people"] = [{
        "name": text_content(el.xpath('.//h3')[0]),
        "category": el.get("data-category"),
        "description": text_content(el.xpath('.//p')[0]),
        "path": el.get("href"),
    } for el in people.xpath('//*[@class="person-card"]')]
    records["teams"] = [text_content(el.xpath('.//h3')[0]) for el in page_soup("teams.html").xpath('//*[@class="team-card"]')]
    records["publication_records"] = len(page_soup("publications.html").xpath('//*[@class="publication-item"]'))
    records["patent_records"] = len(page_soup("achievements.html").xpath('//*[@data-category="patent"]//*[@class="achievement-card"]'))
    records["software_records"] = len(page_soup("achievements.html").xpath('//*[@data-category="software"]//*[@class="achievement-card"]'))
    records["award_records"] = len(page_soup("achievements.html").xpath('//*[@data-category="award"]//*[@class="achievement-card"]'))
    records["video_records"] = len(page_soup("videos.html").xpath('//video|//iframe'))
    AUDIT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    return records


def set_font(style, size, east_asia="STSong", bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in tuple(fonts.attrib):
        if "Theme" in key:
            del fonts.attrib[key]
    fonts.set(qn("w:eastAsia"), east_asia)


def para(doc, text, bold_lead=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(label + "：").bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    return p


def table(doc, headers, rows, widths, size=10.5, padding=95):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, width in zip(t.rows[0].cells, widths):
        cell.width = Inches(width)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value, width in zip(t.add_row().cells, row, widths):
            cell.width = Inches(width)
            cell.text = value
    for ri, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            tr_pr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EEF3" if ri == 0 else ("F7F8FA" if ri % 2 == 0 else "FFFFFF"))
            tc_pr.append(shade)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                b = OxmlElement("w:" + edge)
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "5")
                b.set(qn("w:color"), "D9D9D9")
                borders.append(b)
            tc_pr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top", padding), ("bottom", padding), ("left", 110), ("right", 110)):
                m = OxmlElement("w:" + edge)
                m.set(qn("w:w"), str(value))
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            tc_pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.keep_with_next = False
                if ci == 0 and len(headers) >= 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(4)
    gap.add_run().font.size = Pt(3)
    return t


def section(doc, title):
    heading = doc.add_heading(title, level=1)
    heading.paragraph_format.page_break_before = True


def build(records):
    doc = Document()
    section_settings = doc.sections[0]
    section_settings.page_width = Inches(8.5)
    section_settings.page_height = Inches(11)
    section_settings.top_margin = Inches(0.65)
    section_settings.bottom_margin = Inches(0.65)
    section_settings.left_margin = Inches(0.7)
    section_settings.right_margin = Inches(0.7)
    section_settings.footer_distance = Inches(0.3)
    set_font(doc.styles["Normal"], 11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    set_font(doc.styles["Title"], 21, "Hiragino Sans GB", True)
    doc.styles["Title"].paragraph_format.space_after = Pt(9)
    set_font(doc.styles["Heading 1"], 16, "Hiragino Sans GB", True)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
    set_font(doc.styles["Heading 2"], 12.5, "Hiragino Sans GB", True)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
    set_font(doc.styles["List Bullet"], 11)
    doc.styles["List Bullet"].paragraph_format.line_spacing = 1.2
    doc.styles["List Bullet"].paragraph_format.left_indent = Inches(0.16)
    doc.styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.12)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.core_properties.title = "网站群建设资料收集说明"
    doc.core_properties.subject = FULL_NAME + "网站内容收集与核实"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "网站建设 资料收集 研究方向 研究团队 人才队伍 科研成果 论文 视频"
    footer = section_settings.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for r in footer.runs:
        r.font.size = Pt(9)
    
    doc.add_paragraph("网站群建设资料收集说明", style="Title")
    para(doc, FULL_NAME)
    para(doc, "整理日期  2026年10月3日", after=9)
    para(doc, "为完善实验室网站群，请协调相关老师和团队补充以下资料，并对现有文字、人员职务和成果数据作统一确认。网站栏目和页面框架已搭建，正式发布仍需完整的人员、团队、成果及论文资料，以及审核定稿的介绍和展示素材。")
    para(doc, "当前网站已设两个研究方向、四个暂拟团队、七位老师个人主页，并迁入15篇新闻和2篇通知。专利、软件著作权和论文均无具体记录；获奖仅有2条简略记录；视频页目前保留页头和页尾。上述数量是网站现有记录数，不代表实验室实际规模或成果总量。")
    doc.add_heading("资料收集总览", level=2)
    table(doc, ["栏目", "当前情况", "本次需要提供", "建议对接"], [
        ["实验室简介", "已有基础介绍\n需要定稿核实", "正式简介、组织架构、建设信息", "实验室办公室"],
        ["研究方向", "2个页面\n已有概括文字", "各方向正式介绍、项目、配图与指标", "方向负责人"],
        ["研究团队", "4个暂拟团队\n无负责人和名册", "正式团队名称、介绍、负责人及成员", "各团队负责人"],
        ["人才队伍", "7个简略主页\n职称和名册待审", "完整名册、简历、照片及个人成果", "各位老师"],
        ["科研成果", "专利与软著为空\n获奖2条简略记录", "专利、获奖、软著台账及证明材料", "成果管理人员"],
        ["论文发表", "0条记录\n已设年份与搜索", "完整论文清单及归属确认", "各团队或老师"],
        ["新闻和通知", "新闻15篇\n通知2篇", "迁移范围、完整档案、新增稿件", "宣传或办公室"],
        ["视频展示", "暂无视频内容", "公开视频、封面、简介及发布确认", "各团队"],
        ["品牌和联系", "已有标识与地址\n联系方式不完整", "标识原文件、正式地址、邮箱和电话", "实验室办公室"],
        ["发布和维护", "待明确管理口径", "发布范围、审核人、收集联系人", "实验室负责人"],
    ], [0.9, 1.6, 3.25, 1.35], 10)
    para(doc, "请优先确认正式研究方向介绍、团队设置、完整人员名册和成果收录范围，再分发各栏目资料收集任务。表内对接对象为建议，具体负责人和反馈时间请由实验室确定。", after=0)

    section(doc, "实验室简介与两个研究方向")
    doc.add_heading("实验室简介需要补充和确认的资料", level=2)
    para(doc, "现有简介写明实验室于2021年5月经文化和旅游部批准成立，依托厦门大学，固定人员来自信息学院、电影学院、建筑与土木工程学院、人文学院等单位。请提供可对外发布的正式定稿，并核实成立时间、依托和参与单位、实验室主任及副主任等表述是否仍准确。")
    bullet(doc, "基础信息", "中英文全称、规范简称、批准成立日期、建设定位、主要任务、研究特色及组织架构；如有学术委员会，请提供现任名单和职责，并确认是否展示。")
    bullet(doc, "建设与条件", "发展沿革、实验场地、主要设备和平台、合作单位、人才培养与社会服务情况；如展示人员、项目或成果数量，应同时提供统计日期及统计范围。")
    bullet(doc, "提交材料", "建议提供首页简介200至300字、详情介绍800至1500字及2至4张真实实验室或工作场景照片。字数为便于页面排版的建议，可按实际材料调整。")
    doc.add_heading("两个研究方向分别提供正式材料", level=2)
    table(doc, ["方向", "当前页面文字涉及的内容", "需要确认的范围"], [
        ["闽台非遗文化\n数字化保护", "多模态采集与建档、传统建筑与工艺数字化、数字资产与知识组织、数字展示与公共文化服务。", "哪些内容属于正式研究任务；各子课题的定位、负责人和代表性成果。"],
        ["闽台非遗文化\n智能处理", "方言语音、表演分析与生成、数字人与沉浸式交互、非遗知识计算与智能创意。", "与数字化保护方向的分工；技术、平台、应用案例的正式名称及进展。"],
    ], [1.55, 2.85, 2.7])
    para(doc, "每个方向建议单独提交一份材料：100至150字概述、600至1000字详细介绍、主要研究问题与内容、负责人和参与团队、3至5项代表项目或成果、2至4张配图及图注、可展示的平台或演示地址。项目请写明正式名称、起止时间、负责人、阶段和公开成果，并明确与两个方向的对应关系。")
    doc.add_heading("现有数据和案例需逐项核实", level=2)
    para(doc, "智能处理页面正在展示2237+小时无标签方言音频、1764+小时有标签方言音频和92%准确率。请确认原始数值、是否保留加号、统计时间、语料范围、准确率对应的任务和测试条件，并提供可公开的依据。MinSpeech、闽音智聆、木偶机器人、唱腔合成、数字人导览等名称、功能、团队归属及当前状态也需确认。")

    section(doc, "研究团队资料")
    para(doc, "现有页面按四类任务暂拟团队，尚无负责人、完整成员名册或正式团队介绍。请先确认实验室实际团队数量、正式名称和排序，再决定保留、合并、调整或新增；两个研究方向不等于只能设置两个团队。")
    table(doc, ["序号", "当前暂拟团队名称", "当前概括内容"], [
        ["1", records["teams"][0], "戏曲、音乐、木偶、动作捕捉、机器人表演、数字人与沉浸式展示。"],
        ["2", records["teams"][1], "方言语料、语音识别与合成、方言大模型和智能交互。"],
        ["3", records["teams"][2], "建筑测绘、三维建模、壁画修复、工艺仿真和数字展示。"],
        ["4", records["teams"][3], "数字资产、知识图谱、智能内容生成及文化创意设计。"],
    ], [0.5, 3.35, 3.25])
    doc.add_heading("每个正式团队需要提供的内容", level=2)
    bullet(doc, "团队身份与成员", "中英文名称、简称、负责人姓名和职务、依托学院、成员名单及排序、主要联系人；如展示学生或合作人员，请明确公开范围。")
    bullet(doc, "研究介绍", "100至150字摘要和500至1000字详细介绍，说明研究目标、技术特色、主要任务、对应研究方向及与其他团队的分工。")
    bullet(doc, "项目与成果", "3至5项代表项目或案例，列明项目名称、负责人、起止时间、阶段、成果及团队贡献；提供可关联的论文、专利、获奖和软件清单。")
    bullet(doc, "图片与展示", "团队合照或工作场景、项目图片、图注、拍摄时间及来源；如有演示系统、数据集或开放资源，请提供正式名称和可公开访问地址。")
    para(doc, "目前团队内容集中在总览页。如需要每个团队拥有独立主页，请同时明确主页名称、团队介绍结构、成员和成果展示范围。建议由各团队负责人确认名称、成员、研究描述及可公开材料后统一提交。")

    section(doc, "人才队伍与老师个人主页")
    para(doc, "人才队伍按教授、副教授、助理教授三类展示。现有七位老师均已建立简略主页，照片位置使用姓氏占位，个人信息主要是一两句工作概述，不能替代完整个人简历。下表仅记录当前网页写法，姓名、职称、单位和实验室职务均请逐人核实。")
    table(doc, ["姓名", "当前分类", "当前页面标注", "重点确认"], [
        ["史晓东", "教授", "实验室主任\n信息学院", "现任职务及完整简介"],
        ["姚俊峰", "教授", "实验室副主任\n电影学院", "现任职务及完整简介"],
        ["王量量", "教授", "实验室副主任\n建筑与土木工程学院", "现任职务及完整简介"],
        ["刘昆宏", "教授", "电影学院", "实验室角色和团队归属"],
        ["邱仲潘", "副教授", "未标明所在学院", "姓名用字、单位和职称"],
        ["陈心怡", "助理教授", "未标明所在学院", "单位、职称和成员身份"],
        ["曾华琳", "助理教授", "实验室秘书\n未标明所在学院", "职称、单位和实验室职务"],
    ], [0.85, 1.0, 2.8, 2.45])
    doc.add_heading("请先提供完整名册", level=2)
    para(doc, "现有新闻还提及韩洁、苏劲松、吴清锋、陈俐燕、曹娟等老师，尚未纳入现有人才页。请依据正式人员名单确认是否收录、所属类别及排序；新闻中的参与身份不能直接作为当前实验室成员身份。七人不是已确认的完整名册。")
    doc.add_heading("每位老师需要提交的主页资料", level=2)
    bullet(doc, "必备信息", "规范姓名及英文拼写、现职称、所在学院或单位、实验室职务、所属团队、研究方向和关键词、200至300字简介、一张清晰肖像照、公开工作邮箱。")
    bullet(doc, "详细介绍", "教育与工作经历、主要研究内容、代表项目、代表论文、专利与获奖、学术兼职；可提供原有官方个人主页或完整简历作为整理来源。")
    bullet(doc, "招生与联系", "如需公开，提供招生方向、培养层次、联系方式、办公室及个人主页地址。请由老师本人确认照片、文字和联系方式的公开版本，并写明资料更新时间。")

    section(doc, "科研成果资料")
    para(doc, "科研成果按专利、获奖、软件著作权分类。当前专利和软件著作权栏目为空，获奖栏目仅有两条简略描述，需要提供正式台账和证明材料。请先确认收录年份范围、与实验室相关的成果归属口径及跨团队重复成果的处理方式。")
    doc.add_heading("现有获奖条目需要核实", level=2)
    table(doc, ["当前记录", "缺少或不确定的信息"], [
        ["2024年“礼敬中华优秀传统文化”相关作品获奖", "正式奖项和作品名称、奖项类别、获奖层级、获奖日期、主办或授奖单位、完成人及实验室贡献。"],
        ["鼓浪屿AR导航项目获“挑战杯”国家级银奖", "正式项目名称、比赛届次及赛道、年份、奖项等级、团队成员与指导老师、证书和公开报道。"],
    ], [2.5, 4.6])
    para(doc, "2025年11月13日新闻另报道第九届新加坡规划师学会规划奖中的铜奖和特别提名奖，尚未作为独立条目录入获奖栏目。请确认是否收录，并提供两项项目的完整获奖信息。媒体报道、入选典型案例和资助立项应区分于获奖，由实验室确定展示位置。")
    doc.add_heading("按成果类别提交统一记录", level=2)
    bullet(doc, "专利", "名称、类型、申请号或专利号、申请日、授权日、当前状态、发明人及顺序、申请人或专利权人、所属团队、成果简介及公开证明。请区分申请中和已授权。")
    bullet(doc, "获奖", "奖项全称、获奖项目或作品、年份和届次、类别及等级、授奖单位、完成人及排名、指导老师、实验室贡献、证书或官方公示。")
    bullet(doc, "软件著作权", "软件全称及简称、版本号、登记号、登记日期、著作权人、开发团队、主要功能、应用情况、软件截图和登记证明；演示或下载仅提供允许公开的版本。")
    doc.add_heading("提交格式与统计要求", level=2)
    para(doc, "建议各类成果分别提供Excel台账，一行一项成果；证明材料另附PDF或图片，并与台账序号对应。每项标注关联研究方向、团队和老师，避免在多处出现不同名称或年份。统计数值需统一按成果项数计算还是按人员参与次数计算，并记录截止日期。涉及不能公开的成果，可在汇总中标注“暂不展示”。")

    section(doc, "论文发表资料")
    para(doc, "论文页已设置年份筛选和关键词搜索，目前没有任何论文条目，年份选项为2021至2026年。请提供经团队或作者确认的完整论文清单，并确定是否纳入2021年以前论文、在线发表论文、已录用论文、预印本或其他成果类型。")
    doc.add_heading("每篇论文至少提供的字段", level=2)
    table(doc, ["字段组", "需要填写的内容"], [
        ["文献信息", "论文完整题目、全部作者及顺序、期刊或会议全称、发表年份、卷期页码或文章号。中文或英文按原文填写。"],
        ["识别与访问", "DOI、出版社或正式发表页面、可公开的全文地址；如无DOI，请注明。另附BibTeX或RIS可减少录入偏差。"],
        ["发表状态", "已正式发表、在线发表、已录用或预印本；注明相应日期。确认按在线年份还是正式出版年份筛选。"],
        ["实验室归属", "实验室参与作者、第一或通讯作者身份、署名单位、关联团队和研究方向；说明纳入实验室成果的依据。"],
        ["可选展示", "摘要、关键词、代表论文标记、项目关联、正式论文配图。如显示分区或等级，请注明评价体系、版本和年份。"],
    ], [1.3, 5.8])
    doc.add_heading("需要统一的收录规则", level=2)
    bullet(doc, "范围", "以哪些作者、单位署名或项目关联为纳入依据；合作成果如何标注；各老师个人主页的代表论文是否从总清单同步。")
    bullet(doc, "时间与排序", "收录起止年份、年份判定方式、同一年内的排序；会议论文后续扩展为期刊论文时如何标注和去重。")
    bullet(doc, "全文与数据", "确认哪些论文PDF、数据或代码可以在网站公开。需要保留外部访问地址的，请提供正式且稳定的入口。")
    doc.add_heading("建议统一提交方式", level=2)
    para(doc, "使用Excel一行一篇论文，不要只提交个人主页截图或截断的引用文本。建议表头为：题目、作者、期刊或会议、年份、卷期页码或文章号、DOI、发表状态、链接、实验室作者、所属团队、研究方向、关键词、备注。可以附已有文献管理导出文件或学校科研系统导出表，再由作者确认。")
    para(doc, "论文、成果和老师主页正式补齐后，需要同步加入全站搜索。目前全站搜索收录栏目和新闻摘要等预设信息，老师个人主页仅收录史晓东、姚俊峰、王量量三人，另外四位老师尚未加入搜索索引；搜索索引也不是所有页面全文的自动采集。")

    section(doc, "新闻动态与通知公告资料")
    para(doc, "当前已建立15篇新闻和2篇通知的本地正文页面，正文包含迁入的文字、图片及部分表格。新闻发布时间覆盖2022年12月15日至2026年7月15日；两篇通知发布时间分别为2024年11月23日和11月29日。请确认这些条目是否构成需要保留的完整档案，并补充未迁移的历史或最新内容。")
    doc.add_heading("已有迁移内容需要确认", level=2)
    bullet(doc, "迁移范围", "提供旧站新闻与通知的完整目录、源稿或导出备份，并标明需保留、补充或不再展示的条目。旧站停用前请完整保存原文、原图、附件和发布时间。")
    bullet(doc, "原文一致性", "对每篇核对标题、日期、署名、全文、图片数量和顺序、图注、表格、附件。首页及列表的摘录与正文需采用同一批准版本，历史报道中的原有职务应保留当时口径。")
    bullet(doc, "缺少的署名", "10篇新闻及2篇通知的作者字段为空，例如2024年12月3日和2023年4月13日新闻。请确认沿用空白、隐藏作者字段，还是补充正式供稿单位；不得自行填写。")
    bullet(doc, "原图与附件", "请提交完整图片原文件、会议通知和程序册附件、回执等下载文件，并说明哪些只在正文展示、哪些允许下载。列表封面请明确选用对应原文中的哪张图。")
    doc.add_heading("历史公告存在需由业务人员核实的细节", level=2)
    para(doc, "2024年11月23日会议通知正文出现参会回执11月15日前发送，后附回执备注又写11月25日前发送；通知落款为11月1日，而页面发布日期为11月23日。请核实是否存在不同版本。该文还保留历史缴费码、联系人和回执表，请确认作为历史档案原样保留的展示方式，不应被误作新的有效通知。")
    doc.add_heading("新增稿件的收集格式", level=2)
    para(doc, "每篇提供标题、栏目、发布日期、作者或供稿单位、完整正文、图片及图注、附件、审核人和首页推荐意见。正文建议提交可编辑Word，图片另放文件夹并按出现顺序编号，不能只提供截图或短摘要。新增新闻内容和公告更新应确定固定供稿联系人。")
    para(doc, "如旧稿存在需要更正的事实或版本差异，请提供经过确认的替换稿及更正说明，再同步更新全文、列表和首页。新闻中出现的科研成果和人员信息，需与对应成果台账、人员名册交叉核对。")

    section(doc, "视频展示与网站图片素材")
    doc.add_heading("视频展示内容", level=2)
    para(doc, "视频页面目前只有页头和页尾，尚无播放器或视频条目。请确认首批是否上线视频内容，或者继续保留空栏目。可以收集实验室介绍、项目演示、非遗采集或保护案例、学术交流及媒体报道视频，由实验室确定首批选片和排序。")
    table(doc, ["材料", "每条视频需要提供的信息"], [
        ["基本信息", "标题、100至200字简介、所属团队或研究方向、拍摄或制作日期、发布日期、时长、排序和推荐封面。"],
        ["播放文件", "可公开的成片文件或允许嵌入的稳定播放地址；提供MP4格式较方便。优先1080p，画面清晰、声音完整。"],
        ["辅助素材", "独立封面原图、字幕文件或字幕文本、讲解人或出镜人员、必要的相关项目和论文链接。"],
        ["来源与公开范围", "制作方或媒体来源、允许发布和嵌入的范围；确认人物、背景音乐、作品画面等已具备相应公开使用条件。"],
    ], [1.25, 5.85])
    doc.add_heading("首页与各栏目的图片", level=2)
    para(doc, "首页目前使用四张非遗主题轮播图片，实验室简介和研究方向页也使用部分主题图片。这些素材可用于页面展示，但不能直接作为真实实验项目或团队成果的证明。请确认首页轮播及公共页头最终采用哪些图片，并提供真实、清晰、来源明确的替换素材。新闻图片继续对应新闻原稿。")
    bullet(doc, "轮播与页头", "每张图片提供主题名称、拍摄地点或项目、图注、原始文件和可公开范围。横向原图建议宽度不少于1920像素，主体周边留出裁切空间。")
    bullet(doc, "团队与个人", "团队合照或工作照、老师肖像照、真实科研设备和成果图；个人照片建议长边不少于1000像素，提供统一或易于统一的背景与构图。")
    bullet(doc, "标识与背景", "原实验室Logo保持现有图形，请补充高清透明PNG或矢量源文件、标准色和可使用的反白版本；页尾厦大校园背景需确认图片来源和采用版本。")
    para(doc, "图片建议提交JPG或PNG原文件，并附图片与页面的对应清单。照片中的人物、项目或文物名称请准确说明；示意图、效果图或生成素材应明确标注其性质，正式科研介绍优先使用真实材料。")

    section(doc, "联系信息与资料汇总方式")
    doc.add_heading("公共信息和正式发布事项", level=2)
    para(doc, "当前页尾仅有“福建省厦门市思明区思明南路422号”和邮编361005，没有实验室电话、公共邮箱或具体办公地点。请确认实际实验室所在校区、楼宇、房间及通讯地址，并提供对外咨询联系人、办公电话、公共邮箱；微信公众号或二维码如需展示，请提供正式素材。")
    para(doc, "相关链接目前为厦门大学、文化和旅游部、信息学院、电影学院、建筑与土木工程学院。请确认链接名单、正式单位名称及是否增加人文学院等关联单位。实验室英文全称、版权署名、最终域名及依托学校已有备案信息的展示要求，请由实验室和学校网站管理部门确认。")
    para(doc, "请确定内容审核人、日常更新联系人、公开信息范围和更新频率。是否增设学术委员会、仪器平台、开放课题、合作交流或招生信息等栏目，属于后续栏目选择，请按实际需要确认；本次可先收集材料，不要求全部新增页面。")
    doc.add_heading("资料如何提交", level=2)
    bullet(doc, "介绍类", "每个方向、每个团队和每位老师分别提交可编辑Word；照片和图注另附。建议材料文件夹按栏目及团队或姓名命名。")
    bullet(doc, "清单类", "人员、专利、获奖、软件著作权和论文分别提供Excel；一行一人或一项成果，证明材料与条目编号对应。")
    bullet(doc, "视频与新闻", "视频成片、封面和字幕单独保存；新闻与公告按标题或日期分文件夹，包含完整稿件、全部图片和附件。")
    para(doc, "各份材料请附提交人、联系方式、资料更新时间、确认人、可公开范围及待补事项。如暂时没有某类成果，请明确填写“暂无”；尚未收集齐请填写“待补充”，不要将空白理解为没有成果。建议优先回传已确认的文字和台账，再补齐高清图片、证明材料及视频。")
    doc.add_heading("汇总反馈表", level=2)
    table(doc, ["栏目或对象", "联系人", "资料状态", "确认人及备注"], [
        ["实验室简介及研究方向", "待填写", "待填写", "待填写"],
        ["各研究团队", "待填写", "待填写", "待填写"],
        ["完整名册及老师主页", "待填写", "待填写", "待填写"],
        ["科研成果及论文", "待填写", "待填写", "待填写"],
        ["新闻公告及视频图片", "待填写", "待填写", "待填写"],
        ["联系方式及发布事项", "待填写", "待填写", "待填写"],
    ], [2.1, 1.05, 1.25, 2.7], 10.5, padding=70)
    para(doc, "统一收集联系人：________________    反馈时间：________________", after=0)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "news": len(records["news"]), "notices": len(records["notices"]), "people": len(records["people"]), "teams": len(records["teams"]), "awards": records["award_records"], "papers": records["publication_records"]}, ensure_ascii=False))


if __name__ == "__main__":
    build(audit_site())
